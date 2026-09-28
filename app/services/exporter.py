from __future__ import annotations

import csv
from datetime import datetime, timezone
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import Application


APPLICATION_HEADERS = [
    "ID",
    "Company",
    "Role",
    "Primary Track",
    "Tags",
    "Location",
    "Work Mode",
    "Employment Type",
    "Status",
    "Stage",
    "Priority",
    "Fit /10",
    "Job URL",
    "Application URL",
    "Source",
    "Date Found",
    "Applied Date",
    "Application Deadline",
    "Follow-up Date",
    "Interview Date",
    "Next Action",
    "Salary Expectation",
    "Salary Notes",
    "CV Version",
    "Contact",
    "Notes",
    "Created",
    "Updated",
]


def _applications(session: Session) -> list[Application]:
    return list(
        session.scalars(
            select(Application)
            .options(
                selectinload(Application.tags),
                selectinload(Application.requirements),
                selectinload(Application.documents),
                selectinload(Application.status_history),
                selectinload(Application.activities),
            )
            .order_by(Application.date_found.desc(), Application.id.desc())
        )
    )


def _application_row(application: Application) -> list[object | None]:
    return [
        application.legacy_id or application.id,
        application.company,
        application.role,
        application.primary_track,
        ", ".join(tag.name for tag in application.tags),
        application.location,
        application.work_mode,
        application.employment_type,
        application.status,
        application.pipeline_stage,
        application.priority,
        application.fit_score,
        application.job_url,
        application.application_url,
        application.source,
        application.date_found,
        application.applied_date,
        application.application_deadline,
        application.follow_up_date,
        application.interview_date,
        application.next_action,
        application.salary_expectation,
        application.salary_notes,
        application.cv_version,
        application.contact,
        application.notes,
        application.created_at,
        application.updated_at,
    ]


def _excel_value(value: object | None) -> object | None:
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def build_csv(session: Session) -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(APPLICATION_HEADERS)
    for application in _applications(session):
        row = _application_row(application)
        writer.writerow(
            [item.isoformat() if hasattr(item, "isoformat") else item for item in row]
        )
    return stream.getvalue().encode("utf-8-sig")


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="255849")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in column[:40]]
        width = min(max(max((len(value) for value in values), default=0) + 2, 11), 45)
        sheet.column_dimensions[column[0].column_letter].width = width


def build_xlsx(session: Session) -> bytes:
    applications = _applications(session)
    workbook = Workbook()
    application_sheet = workbook.active
    application_sheet.title = "Applications"
    application_sheet.append(APPLICATION_HEADERS)
    for application in applications:
        application_sheet.append([_excel_value(item) for item in _application_row(application)])
    _style_sheet(application_sheet)

    requirements = workbook.create_sheet("Requirements")
    requirements.append(
        ["Application ID", "Company", "Name", "Category", "Required", "Completed", "Value", "Notes", "Sort Order"]
    )
    for application in applications:
        for item in application.requirements:
            requirements.append(
                [application.legacy_id or application.id, application.company, item.name, item.category, item.required, item.completed, item.value, item.notes, item.sort_order]
            )
    _style_sheet(requirements)

    documents = workbook.create_sheet("Documents")
    documents.append(
        ["Application ID", "Company", "Document Type", "Version", "Local Path", "Cloud URL", "Notes"]
    )
    for application in applications:
        for document in application.documents:
            documents.append(
                [application.legacy_id or application.id, application.company, document.document_type, document.version, document.local_path, document.cloud_url, document.notes]
            )
    _style_sheet(documents)

    history = workbook.create_sheet("Status History")
    history.append(
        ["Application ID", "Company", "From Status", "To Status", "Changed At", "Notes"]
    )
    for application in applications:
        for event in reversed(application.status_history):
            history.append(
                [application.legacy_id or application.id, application.company, event.from_status, event.to_status, _excel_value(event.changed_at), event.notes]
            )
    _style_sheet(history)

    activities = workbook.create_sheet("Activities")
    activities.append(
        [
            "Application ID",
            "Company",
            "Occurred At",
            "Type",
            "Direction",
            "Summary",
            "Interview Stage",
            "Scheduled At",
            "Interview Format",
            "Contact",
            "Meeting URL",
            "Needs Action",
            "Action Completed",
            "Notes",
        ]
    )
    for application in applications:
        for activity in reversed(application.activities):
            activities.append(
                [
                    application.legacy_id or application.id,
                    application.company,
                    _excel_value(activity.occurred_at),
                    activity.activity_type,
                    activity.direction,
                    activity.summary,
                    activity.interview_stage,
                    _excel_value(activity.scheduled_at),
                    activity.interview_format,
                    activity.contact,
                    activity.meeting_url,
                    activity.needs_action,
                    activity.action_completed,
                    activity.notes,
                ]
            )
    _style_sheet(activities)

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
