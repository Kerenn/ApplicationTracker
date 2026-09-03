from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.constants import ApplicationStatus, PrimaryTrack, enum_values


DIRECT_MAPPING = {
    "App ID": "legacy_id",
    "Company": "company",
    "Role": "role",
    "Location": "location",
    "Work Mode": "work_mode",
    "Priority": "priority",
    "Fit /10": "fit_score",
    "Job URL": "job_url",
    "Source": "source",
    "Date Found": "date_found",
    "Applied Date": "applied_date",
    "Application Deadline": "application_deadline",
    "Follow-up Date": "follow_up_date",
    "Next Action": "next_action",
    "Salary Expectation (€ gross/year)": "salary_expectation",
    "Salary Notes": "salary_notes",
    "CV Version": "cv_version",
    "Recruiter / Contact": "contact",
    "Interview Date": "interview_date",
    "Last Updated": "updated_at",
}

AMBIGUOUS_HEADERS = {
    "Track": "Combined legacy categories need row-by-row track selection.",
    "Cover Letter?": "May mean required, available, or complete.",
    "Outcome": "May be a status transition or a free-form note.",
    "Fit / Why Apply": "Recommended: append as a labeled Notes section.",
    "Gaps / Tailoring": "Recommended: append as a labeled Notes section.",
    "CV File Path": "Recommended: create a CV Document record.",
    "CV Link": "Recommended: attach to the CV Document record.",
    "Cover Letter File Path": "Recommended: create a Cover Letter Document record.",
    "Cover Letter Link": "Recommended: attach to the Cover Letter Document record.",
}

CANONICAL_TRACKS = set(enum_values(PrimaryTrack))
STATUS_MAPPING = {
    "Saved": (ApplicationStatus.SAVED.value, "Saved"),
    "Preparing": (ApplicationStatus.PREPARING.value, "Reviewing role"),
    "Applying": (ApplicationStatus.APPLYING.value, "Application form"),
    "Applied": (ApplicationStatus.APPLIED.value, "Submitted"),
    "Recruiter Screen": (ApplicationStatus.INTERVIEW.value, "Recruiter screen"),
    "Interview 1": (ApplicationStatus.INTERVIEW.value, "Hiring manager interview"),
    "Technical Interview": (ApplicationStatus.INTERVIEW.value, "Technical interview"),
    "Final Interview": (ApplicationStatus.INTERVIEW.value, "Final interview"),
    "Offer": (ApplicationStatus.OFFER.value, "Offer received"),
    "Rejected": (ApplicationStatus.REJECTED.value, "Rejected"),
    "Withdrawn": (ApplicationStatus.WITHDRAWN.value, "Withdrawn"),
}

IGNORED_HEADERS = {
    "Days Since Applied": "Derived value; calculate dynamically.",
}


@dataclass
class PreviewRow:
    source_row: int
    direct: dict[str, Any] = field(default_factory=dict)
    unresolved: dict[str, Any] = field(default_factory=dict)


@dataclass
class WorkbookPreview:
    path: str
    sheet_names: list[str]
    application_headers: list[str]
    rows: list[PreviewRow]
    cover_letter_rows: int
    warnings: list[str]


def _display(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def preview_workbook(path_value: str | Path, row_limit: int = 100) -> WorkbookPreview:
    path = Path(path_value).expanduser()
    if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("Select an .xlsx or .xlsm workbook.")
    if not path.is_file():
        raise ValueError("Workbook not found at that local path.")

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_vba=path.suffix.casefold() == ".xlsm",
    )
    try:
        if "Applications" not in workbook.sheetnames:
            raise ValueError("The workbook does not contain an Applications sheet.")
        sheet = workbook["Applications"]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        rows: list[PreviewRow] = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, "") for value in cells):
                continue
            source = dict(zip(headers, cells, strict=False))
            preview = PreviewRow(source_row=row_number)
            for header, target in DIRECT_MAPPING.items():
                if header in source and source[header] not in (None, ""):
                    preview.direct[target] = _display(source[header])
            raw_track = source.get("Track")
            if raw_track in CANONICAL_TRACKS:
                preview.direct["primary_track"] = raw_track
            elif raw_track not in (None, ""):
                preview.unresolved["Track"] = _display(raw_track)
            raw_status = source.get("Status")
            if raw_status in STATUS_MAPPING:
                mapped_status, mapped_stage = STATUS_MAPPING[raw_status]
                preview.direct["status"] = mapped_status
                preview.direct["pipeline_stage"] = mapped_stage
            elif raw_status not in (None, ""):
                preview.unresolved["Status"] = _display(raw_status)
            for header in AMBIGUOUS_HEADERS:
                if header == "Track":
                    continue
                if header in source and source[header] not in (None, ""):
                    preview.unresolved[header] = _display(source[header])
            rows.append(preview)
            if len(rows) >= row_limit:
                break

        cover_letter_rows = 0
        if "Cover Letters" in workbook.sheetnames:
            cover_sheet = workbook["Cover Letters"]
            cover_letter_rows = sum(
                1
                for row in cover_sheet.iter_rows(min_row=2, values_only=True)
                if any(value not in (None, "") for value in row)
            )
        warnings = [
            f"{header}: {reason}" for header, reason in AMBIGUOUS_HEADERS.items() if header in headers
        ]
        if any("Status" in row.unresolved for row in rows):
            warnings.append("Status: One or more values need canonical-status confirmation.")
        warnings.extend(
            f"{header}: {reason}" for header, reason in IGNORED_HEADERS.items() if header in headers
        )
        return WorkbookPreview(
            path=str(path),
            sheet_names=list(workbook.sheetnames),
            application_headers=headers,
            rows=rows,
            cover_letter_rows=cover_letter_rows,
            warnings=warnings,
        )
    finally:
        vba_archive = getattr(workbook, "vba_archive", None)
        if vba_archive is not None:
            vba_archive.close()
        workbook.close()
