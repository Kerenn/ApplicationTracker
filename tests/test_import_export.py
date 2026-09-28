from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.models import Application, ApplicationActivity
from app.services.applications import apply_application_form
from app.services.exporter import build_csv, build_xlsx
from app.services.workbook_importer import preview_workbook
from tests.test_applications import application_form


def test_exports_are_readable(session):
    application = apply_application_form(session, Application(), application_form())
    session.add(
        ApplicationActivity(
            application=application,
            activity_type="Interview",
            direction="Internal",
            summary="Technical interview scheduled",
            interview_stage="Technical interview",
            interview_format="Video",
            contact="Alex Recruiter",
            meeting_url="https://meet.example/interview",
        )
    )
    session.commit()

    csv_data = build_csv(session).decode("utf-8-sig")
    assert "Example Robotics" in csv_data
    assert "Computer Vision" in csv_data

    workbook = load_workbook(BytesIO(build_xlsx(session)))
    assert workbook.sheetnames == ["Applications", "Requirements", "Documents", "Status History", "Activities"]
    assert workbook["Applications"]["B2"].value == "Example Robotics"
    assert workbook["Activities"]["J2"].value == "Alex Recruiter"
    assert workbook["Activities"]["K2"].value == "https://meet.example/interview"
    workbook.close()


def test_workbook_preview_does_not_import_and_flags_ambiguity(tmp_path):
    path = tmp_path / "source.xlsm"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"
    sheet.append(["App ID", "Company", "Role", "Track", "Status", "Days Since Applied"])
    sheet.append(["APP-001", "Example", "Engineer", "Embedded/Electronics", "Recruiter Screen", 10])
    workbook.create_sheet("Cover Letters").append(["Application ID", "Company"])
    workbook.save(path)

    preview = preview_workbook(path)
    assert preview.rows[0].direct["company"] == "Example"
    assert preview.rows[0].unresolved["Track"] == "Embedded/Electronics"
    assert preview.rows[0].direct["status"] == "Interview"
    assert preview.rows[0].direct["pipeline_stage"] == "Recruiter screen"
    assert any("Days Since Applied" in warning for warning in preview.warnings)
